import os
from pathlib import Path

from pydantic import BaseModel, Field

from files.indexer import extract_text as extract_document_text
from files.indexer import search as search_index
from memory.activity import get_recent, record_path
from memory.store import get_allowed_dirs

# Extensions edit_document / read_code_file / write_code_file will touch.
# Kept narrow on purpose — this is "help edit text/code," not "run arbitrary
# scripts I write."
TEXT_EDITABLE_EXTS = {".txt", ".md"}
CODE_EXTS = {
    ".py", ".js", ".ts", ".tsx", ".jsx", ".json", ".css", ".html", ".c", ".cpp",
    ".h", ".java", ".cs", ".go", ".rs", ".sql", ".yaml", ".yml", ".sh", ".toml",
}
DOCUMENT_READABLE_EXTS = {".txt", ".md", ".pdf", ".docx"}


def is_within_allowed(path: str) -> bool:
    try:
        resolved = Path(path).resolve()
    except OSError:
        return False
    for root in get_allowed_dirs():
        try:
            resolved.relative_to(Path(root).resolve())
            return True
        except ValueError:
            continue
    return False


class SearchFilesArgs(BaseModel):
    query: str = Field(..., description="Natural-language description of the file to find, e.g. 'DBMS project pdf'")


def search_files(args: SearchFilesArgs) -> dict:
    results = search_index(args.query)
    if not results:
        return {"status": "ok", "count": 0, "results": [], "message": "No matching files found."}
    return {"status": "ok", "count": len(results), "results": results}


class OpenFileArgs(BaseModel):
    path: str = Field(..., description="Full file path to open, usually one returned by search_files")


def open_file(args: OpenFileArgs) -> dict:
    if not is_within_allowed(args.path):
        return {"status": "error", "message": "That path is outside NOVA's allowed folders."}
    if not Path(args.path).exists():
        return {"status": "error", "message": "That file doesn't exist anymore."}
    os.startfile(args.path)  # noqa: S606 — path validated against the allowlist above
    record_path(args.path)
    return {"status": "ok", "message": f"Opened {Path(args.path).name}."}


class OpenFolderArgs(BaseModel):
    path: str = Field(..., description="Full folder path to open in the file explorer")


def open_folder(args: OpenFolderArgs) -> dict:
    if not is_within_allowed(args.path):
        return {"status": "error", "message": "That path is outside NOVA's allowed folders."}
    if not Path(args.path).is_dir():
        return {"status": "error", "message": "That folder doesn't exist."}
    os.startfile(args.path)  # noqa: S606
    return {"status": "ok", "message": f"Opened {args.path} in the file explorer."}


class CreateFolderArgs(BaseModel):
    name: str = Field(..., description="Name of the new folder, e.g. 'TechCommons Hackathon'")
    location: str | None = Field(
        None,
        description="Which allowed folder to create it in, by name (e.g. 'Desktop', 'Documents'). "
        "Defaults to Desktop if not specified.",
    )


def _resolve_location(location: str | None) -> Path:
    dirs = get_allowed_dirs()
    if location:
        for d in dirs:
            if Path(d).name.lower() == location.strip().lower():
                return Path(d)
    for d in dirs:
        if Path(d).name.lower() == "desktop":
            return Path(d)
    return Path(dirs[0])


def create_folder(args: CreateFolderArgs) -> dict:
    base = _resolve_location(args.location)
    target = base / args.name
    if not is_within_allowed(str(target)):
        return {"status": "error", "message": "That path is outside NOVA's allowed folders."}
    target.mkdir(parents=True, exist_ok=True)
    return {"status": "ok", "message": f"Created folder '{args.name}' in {base.name}.", "path": str(target)}


class RecentFilesArgs(BaseModel):
    pass


def recent_files(_: RecentFilesArgs) -> dict:
    """Files the user recently copied (via Explorer, tracked from the
    clipboard) or that NOVA recently opened — lets 'where did I put the
    thing I just copied' work without a fresh full-text search."""
    entries = get_recent(limit=10)
    if not entries:
        return {"status": "ok", "count": 0, "results": [], "message": "Nothing recent tracked yet."}
    results = [{"path": e["path"], "filename": Path(e["path"]).name} for e in entries]
    return {"status": "ok", "count": len(results), "results": results}


class EditDocumentArgs(BaseModel):
    path: str = Field(..., description="Full path to an existing .txt, .md, or .docx file")
    text: str = Field(..., description="Text to add to the document")
    mode: str = Field("append", description="'append' to add to the end, or 'replace' to overwrite the whole file")


def edit_document(args: EditDocumentArgs) -> dict:
    if not is_within_allowed(args.path):
        return {"status": "error", "message": "That path is outside NOVA's allowed folders."}
    path = Path(args.path)
    ext = path.suffix.lower()

    if ext in TEXT_EDITABLE_EXTS:
        if args.mode == "replace" or not path.exists():
            path.write_text(args.text, encoding="utf-8")
        else:
            with path.open("a", encoding="utf-8") as f:
                f.write("\n" + args.text)
        return {"status": "ok", "message": f"Updated {path.name}."}

    if ext == ".docx":
        import docx

        doc = docx.Document(str(path)) if path.exists() else docx.Document()
        if args.mode == "replace":
            for p in list(doc.paragraphs):
                p.text = ""
        doc.add_paragraph(args.text)
        doc.save(str(path))
        return {"status": "ok", "message": f"Updated {path.name}."}

    return {"status": "error", "message": f"I can only edit .txt, .md, or .docx files, not {ext}."}


class ReadDocumentArgs(BaseModel):
    path: str = Field(..., description="Full path to a .txt, .md, .pdf, or .docx file, usually from search_files")


def read_document(args: ReadDocumentArgs) -> dict:
    if not is_within_allowed(args.path):
        return {"status": "error", "message": "That path is outside NOVA's allowed folders."}
    path = Path(args.path)
    if path.suffix.lower() not in DOCUMENT_READABLE_EXTS:
        return {"status": "error", "message": f"{path.suffix} isn't a document type I can read (.txt/.md/.pdf/.docx only)."}
    if not path.exists():
        return {"status": "error", "message": "That file doesn't exist."}
    text = extract_document_text(path)
    if not text.strip():
        return {"status": "error", "message": "Couldn't extract any text from that file — it may be scanned/image-based."}
    return {"status": "ok", "content": text}


class ReadCodeFileArgs(BaseModel):
    path: str = Field(..., description="Full path to a source code file to read")


def read_code_file(args: ReadCodeFileArgs) -> dict:
    if not is_within_allowed(args.path):
        return {"status": "error", "message": "That path is outside NOVA's allowed folders."}
    path = Path(args.path)
    if path.suffix.lower() not in CODE_EXTS:
        return {"status": "error", "message": f"{path.suffix} isn't a recognized code file type."}
    if not path.exists():
        return {"status": "error", "message": "That file doesn't exist."}
    return {"status": "ok", "content": path.read_text(encoding="utf-8", errors="ignore")[:8000]}


class WriteCodeFileArgs(BaseModel):
    filename: str = Field(..., description="File name to write, e.g. 'hello.py'")
    content: str = Field(..., description="The complete new content of the file")
    location: str | None = Field(
        None,
        description="Full folder path to write into, OR a friendly allowed-folder name (e.g. 'Desktop'). "
        "Defaults to Desktop if not specified.",
    )


def write_code_file(args: WriteCodeFileArgs) -> dict:
    base = Path(args.location) if args.location and Path(args.location).is_absolute() else _resolve_location(args.location)
    path = base / args.filename
    if not is_within_allowed(str(path)):
        return {"status": "error", "message": "That path is outside NOVA's allowed folders."}
    if path.suffix.lower() not in CODE_EXTS:
        return {"status": "error", "message": f"{path.suffix} isn't a recognized code file type."}
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(args.content, encoding="utf-8")
    return {"status": "ok", "message": f"Wrote {path.name} to {base.name}.", "path": str(path)}
